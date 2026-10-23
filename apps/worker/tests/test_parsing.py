import io

import pytest
from docx import Document as DocxDocument

from daleel_worker.parsing import (
    DOCX,
    PDF,
    UnsupportedDocumentError,
    clean,
    parse,
    parse_docx,
    parse_pdf,
)
from tests.conftest import DEMO_DIR


def test_arabic_pdf_reads_in_logical_order() -> None:
    parsed = parse_pdf((DEMO_DIR / "procurement-policy-ar.pdf").read_bytes())
    texts = [b.text for b in parsed.blocks]
    headings = [b.text for b in parsed.blocks if b.kind == "heading"]

    assert parsed.page_count == 2
    assert headings[0] == "سياسة المشتريات والتعاقد مع الموردين"
    assert "3. طلب العروض والمنافسة" in headings
    # Word order, ligatures (المشتريات, التي) and numbers inside Arabic all come out right.
    assert (
        "العقود التي تتجاوز 500,000 درهم: تعرض على لجنة المشتريات، ويعتمدها الرئيس التنفيذي."
        in texts
    )
    assert any("الإصدار 2.1" in t for t in texts)
    # Chromium drops an alef written just before a lam-alef ligature from the text layer; the
    # hole its glyph leaves brings it back.
    assert "2. حدود صلاحيات الاعتماد" in headings
    assert any("دون حد الاعتماد،" in t for t in texts)


def test_pdf_blocks_carry_page_and_box() -> None:
    parsed = parse_pdf((DEMO_DIR / "procurement-policy-ar.pdf").read_bytes())
    registration = next(b for b in parsed.blocks if b.text == "4. تسجيل الموردين")
    assert registration.page == 2
    box = registration.boxes[0]
    assert box.page == 2
    assert 0 < box.x0 < box.x1 <= 595  # A4 width in points
    assert 0 < box.y0 < box.y1 <= 842


def test_english_pdf_headings_and_text() -> None:
    parsed = parse_pdf((DEMO_DIR / "information-security-policy-en.pdf").read_bytes())
    headings = [b.text for b in parsed.blocks if b.kind == "heading"]
    assert headings[:3] == ["Information Security Policy", "1. Purpose", "2. Data classification"]
    assert any("at least 14 characters" in b.text for b in parsed.blocks)


def test_docx_headings_paragraphs_and_tables() -> None:
    document = DocxDocument()
    document.add_heading("Travel Policy", level=1)
    document.add_paragraph("Economy class for flights under 6 hours.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Per diem"
    table.rows[0].cells[1].text = "400 AED"
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse_docx(buffer.getvalue())
    assert [(b.kind, b.text) for b in parsed.blocks] == [
        ("heading", "Travel Policy"),
        ("body", "Economy class for flights under 6 hours."),
        ("body", "Per diem | 400 AED"),
    ]
    assert parsed.page_count is None
    assert parse(buffer.getvalue(), DOCX).blocks == parsed.blocks


def test_unsupported_type() -> None:
    with pytest.raises(UnsupportedDocumentError):
        parse(b"hello", "text/plain")
    assert PDF == "application/pdf"


def test_clean_folds_presentation_forms_and_spaces() -> None:
    assert clean("  ﻣﺮﺣﺒﺎ \n  world ") == "مرحبا world"
