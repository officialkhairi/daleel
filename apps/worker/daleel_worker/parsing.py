"""Turns uploaded files into text blocks that remember where they came from (page + box), so a
citation can later point at the exact passage."""

import io
import re
import statistics
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Literal

import pymupdf
from docx import Document as DocxDocument

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SUPPORTED_TYPES = (PDF, DOCX)

BlockKind = Literal["heading", "body"]


@dataclass(frozen=True)
class Box:
    page: int  # 1-based
    x0: float
    y0: float
    x1: float
    y1: float

    def as_dict(self) -> dict[str, float | int]:
        return {"page": self.page, "x0": self.x0, "y0": self.y0, "x1": self.x1, "y1": self.y1}


@dataclass
class Block:
    text: str
    kind: BlockKind
    page: int | None = None
    boxes: list[Box] = field(default_factory=list)


@dataclass
class ParsedDocument:
    blocks: list[Block]
    page_count: int | None


class UnsupportedDocumentError(Exception):
    pass


_SPACES = re.compile(r"\s+")
_ARABIC = re.compile(r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]")
_ARABIC_MARK = re.compile(r"[\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED\u0640]")  # + tatweel
_LATIN_OR_DIGIT = re.compile(r"[A-Za-z0-9]")
# A left-to-right run inside Arabic text: Latin words and numbers, with the punctuation and
# spaces between them ("50,000", "IBM Plex", "2.1").
_LTR_RUN = re.compile(r"[A-Za-z0-9](?:[A-Za-z0-9.,:/%'\- ]*[A-Za-z0-9])?")
_PAGE_NUMBER = re.compile(r"^(page\s*)?\d+(\s*(/|of|من)\s*\d+)?$", re.IGNORECASE)


def clean(text: str) -> str:
    """NFKC folds Arabic presentation forms (what many PDFs store) back into ordinary letters."""
    return _SPACES.sub(" ", unicodedata.normalize("NFKC", text)).strip()


def _line_text(line: dict[str, Any]) -> str:
    """Text of one line in logical (reading) order.

    PDF text order isn't reliable for Arabic: generators emit glyphs in visual order, and
    ligatures map back to their letters reversed. Each character still has its own box, so for a
    mostly-Arabic line we read characters right to left by position, then put embedded
    left-to-right runs (numbers, Latin words) back in their own order.
    """
    chars = [c for span in line["spans"] for c in span["chars"]]
    text = "".join(c["c"] for c in chars)
    if len(_ARABIC.findall(text)) <= len(_LATIN_OR_DIGIT.findall(text)):
        return text

    # A ligature glyph (e.g. lam-alef) yields several characters: the extra ones get a zero-width
    # box at the glyph's right edge and precede, in file order, the character holding the glyph's
    # box. Logically they follow it, so pair them up before ordering by position.
    units: list[tuple[float, float, str]] = []  # (x0, x1, text)
    i = 0
    while i < len(chars):
        char = chars[i]
        x0, _, x1, _ = char["bbox"]
        partner = chars[i + 1] if i + 1 < len(chars) else None
        if x1 - x0 < 0.01 and partner is not None and abs(partner["bbox"][2] - x0) < 0.5:
            units.append((partner["bbox"][0], x0, partner["c"] + char["c"]))
            i += 2
            continue
        units.append((x0, x1, char["c"]))
        i += 1

    # The file's own spaces can't be trusted to sit between words either; like MuPDF does for
    # ordinary text, a space is wherever there is a visible gap between neighbours.
    size = max((span["size"] for span in line["spans"]), default=10.0)
    letters = [u for u in units if not u[2].isspace() and not _ARABIC_MARK.fullmatch(u[2])]
    letters.sort(key=lambda u: -(u[0] + u[1]))
    parts: list[str] = []
    for index, (x0, x1, char_text) in enumerate(letters):
        if index:
            gap = letters[index - 1][0] - x1
            # A zero-width unit lost its glyph's box, which leaves a hole of up to one letter
            # just before it; only a wider gap is a real space.
            threshold = size * (0.45 if x1 - x0 < 0.01 else 0.15)
            if gap > threshold:
                parts.append(" ")
        parts.append(char_text)
    # Diacritics carry boxes of their own and land in odd places; search ignores them anyway
    # (the PDF itself still shows them).
    rtl = _ARABIC_MARK.sub("", "".join(parts))
    return _LTR_RUN.sub(lambda m: m.group(0)[::-1], rtl)


def _is_bold(span: dict[str, Any]) -> bool:
    return bool(span["flags"] & 16) or "bold" in span.get("font", "").lower()


def parse_pdf(data: bytes) -> ParsedDocument:
    with pymupdf.open(stream=data, filetype="pdf") as doc:  # type: ignore[no-untyped-call]
        raw: list[tuple[int, dict[str, Any]]] = []
        sizes: list[float] = []
        for page_index, page in enumerate(doc, start=1):
            for block in page.get_text("rawdict", sort=True)["blocks"]:
                if block.get("type") != 0:
                    continue  # images
                raw.append((page_index, block))
                for line in block["lines"]:
                    for span in line["spans"]:
                        visible = sum(1 for c in span["chars"] if not c["c"].isspace())
                        sizes.extend([span["size"]] * visible)
        page_count = doc.page_count

    body_size = statistics.median(sizes) if sizes else 0.0
    blocks: list[Block] = []
    for page_number, block in raw:
        spans = [
            s
            for line in block["lines"]
            for s in line["spans"]
            if any(not c["c"].isspace() for c in s["chars"])
        ]
        text = clean(" ".join(_line_text(line) for line in block["lines"]))
        if not text or _PAGE_NUMBER.match(text):
            continue
        largest = max(s["size"] for s in spans)
        all_bold = all(_is_bold(s) for s in spans)
        is_heading = largest >= body_size * 1.15 or (
            all_bold and len(text) <= 120 and not text.endswith((".", "。", ":"))
        )
        x0, y0, x1, y1 = block["bbox"]
        blocks.append(
            Block(
                text=text,
                kind="heading" if is_heading else "body",
                page=page_number,
                boxes=[Box(page_number, round(x0, 1), round(y0, 1), round(x1, 1), round(y1, 1))],
            )
        )
    return ParsedDocument(blocks=blocks, page_count=page_count)


def parse_docx(data: bytes) -> ParsedDocument:
    document = DocxDocument(io.BytesIO(data))
    blocks: list[Block] = []
    for paragraph in document.paragraphs:
        text = clean(paragraph.text)
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        kind: BlockKind = "heading" if style.startswith(("Heading", "Title")) else "body"
        blocks.append(Block(text=text, kind=kind))
    for table in document.tables:
        for row in table.rows:
            text = clean(" | ".join(cell.text for cell in row.cells))
            if text:
                blocks.append(Block(text=text, kind="body"))
    return ParsedDocument(blocks=blocks, page_count=None)


def parse(data: bytes, content_type: str) -> ParsedDocument:
    if content_type == PDF:
        return parse_pdf(data)
    if content_type == DOCX:
        return parse_docx(data)
    raise UnsupportedDocumentError(f"Unsupported file type: {content_type}")
